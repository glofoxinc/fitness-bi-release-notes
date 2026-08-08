#!/usr/bin/env python3
"""Generate Fitness BI release notes Word doc from tickets JSON."""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt


MONTHS = [
    "",
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]


def parse_release_date(fix_version: str) -> str:
    # Version format: YYYY.Quarter.MM.DD  (e.g. 2026.3.08.05 -> 05-August-2026)
    parts = fix_version.strip().split(".")
    if len(parts) >= 4:
        year, month, day = int(parts[0]), int(parts[2]), int(parts[3])
    elif len(parts) == 3:
        # Fallback YYYY.MM.DD
        year, month, day = int(parts[0]), int(parts[1]), int(parts[2])
    else:
        raise ValueError(f"Unexpected fix version format: {fix_version}")
    return f"{day:02d}-{MONTHS[month]}-{year}"


FONT_NAME = "Cambria"
BODY_SIZE = 12
HEADING_SIZE = 14
TITLE_SIZE = 16
TABLE_SIZE = 10


def set_run_font(run, bold=False, size=BODY_SIZE):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FONT_NAME
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), FONT_NAME)
    rFonts.set(qn("w:hAnsi"), FONT_NAME)
    rFonts.set(qn("w:cs"), FONT_NAME)


def add_heading_line(doc: Document, text: str, size=TITLE_SIZE, bold=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, bold=bold, size=size)
    return p


def add_kv(doc: Document, label: str, value: str):
    p = doc.add_paragraph()
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, bold=True, size=BODY_SIZE)
    r2 = p.add_run(value)
    set_run_font(r2, bold=False, size=BODY_SIZE)
    return p


def add_section_title(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=True, size=HEADING_SIZE)
    return p


def add_body(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=False, size=BODY_SIZE)
    return p


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, bold=False, size=BODY_SIZE)


def set_cell_text(cell, text: str, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text if text is not None else "")
    set_run_font(run, bold=bold, size=TABLE_SIZE)


def set_cell_hyperlink(cell, url: str):
    """Write a real clickable URL instead of plain text."""
    cell.text = ""
    if not url:
        return

    paragraph = cell.paragraphs[0]
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_fonts = OxmlElement("w:rFonts")
    run_fonts.set(qn("w:ascii"), FONT_NAME)
    run_fonts.set(qn("w:hAnsi"), FONT_NAME)
    run_properties.append(run_fonts)

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(TABLE_SIZE * 2))
    run_properties.append(size)

    text = OxmlElement("w:t")
    text.text = url
    run.append(run_properties)
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def shade_header_row(row):
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9E2F3")
        shd.set(qn("w:val"), "clear")
        tc_pr.append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, bold=True)
    shade_header_row(hdr)
    for r_idx, row_vals in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_vals):
            set_cell_text(row.cells[c_idx], val)
    return table


def distinct_clients(tickets: list[dict], typ: str) -> list[str]:
    seen = []
    for t in tickets:
        if t.get("type") != typ:
            continue
        c = t.get("client") or ""
        if c and c not in seen:
            seen.append(c)
    return seen


def build_document(data: dict) -> Document:
    version = data["fix_version"]
    meta = data.get("meta") or {}
    narratives = data.get("narratives") or {}
    tickets = data.get("tickets") or []
    excluded = data.get("excluded") or []
    mode = data.get("mode", "pre")

    release_date = data.get("release_date") or parse_release_date(version)
    analyze_clients = distinct_clients(tickets, "Analyze")
    custom_clients = distinct_clients(tickets, "Custom")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    add_heading_line(doc, "Release Notes", size=18)
    add_kv(doc, "Version", f"v{version}")
    add_kv(doc, "Release Date", release_date)
    add_kv(doc, "Development Team", meta.get("development_team", ""))
    add_kv(doc, "Release Owner", meta.get("release_owner", ""))
    add_kv(doc, "Senior Manager", meta.get("senior_manager", ""))
    add_kv(doc, "Director", meta.get("director", ""))
    add_kv(doc, "SVP", meta.get("svp", ""))
    add_kv(doc, "Environment", meta.get("environment", "Production"))

    add_section_title(doc, "Release Summary")
    add_body(
        doc,
        narratives.get("release_summary")
        or "This release includes reporting enhancements, custom client updates, and semantic model improvements.",
    )

    add_section_title(doc, "Key Highlights")
    highlights = []
    for t in tickets:
        # Key Highlights shows concise ticket info and is independent of the
        # name-only Report column. Falls back to Client: Report only if needed.
        line = t.get("highlight")
        if not line:
            client = t.get("client") or ""
            report = t.get("report") or t.get("description") or ""
            line = f"{client}: {report}" if client and report else report
        if line:
            highlights.append(line)
    add_bullets(doc, highlights or ["(No tickets found)"])

    add_section_title(doc, "Affected Clients:")
    add_body(doc, "Analyze Clients")
    if analyze_clients:
        add_bullets(doc, ["All Analyze-based (Standard) clients"])
    else:
        add_bullets(doc, ["None in this release"])

    add_section_title(doc, "Custom Clients Summary:")
    add_body(
        doc,
        narratives.get("custom_clients_summary")
        or (
            "This release includes updates to existing functionality, semantic model changes, "
            "column additions, and new report creation for custom clients."
            if custom_clients
            else "No custom client changes in this release."
        ),
    )
    if custom_clients:
        add_bullets(doc, custom_clients)

    add_section_title(doc, "Standard & Custom Summary:")
    add_body(
        doc,
        narratives.get("standard_custom_summary")
        or (
            "This release improves reporting usability through dashboard restructuring, clearer "
            "report metrics, and standardized paginated report formatting alongside custom client enhancements."
        ),
    )

    add_section_title(doc, "Dashboards / Reports:")
    dash = []
    for t in tickets:
        client = t.get("client") or ""
        report = t.get("report") or ""
        if client and report:
            dash.append(f"{client} - {report}")
        elif report:
            dash.append(report)
    add_bullets(doc, dash or ["(None)"])

    add_section_title(doc, "Deployment Details:")
    deploy_clients = []
    if analyze_clients:
        deploy_clients.append("Analyze")
    for c in custom_clients:
        if c not in deploy_clients:
            deploy_clients.append(c)
    deploy_rows = []
    for c in deploy_clients:
        # Pre-deploy: leave owner/status/id blank for the team
        deploy_rows.append([c, "", "", ""])
    if not deploy_rows:
        deploy_rows.append(["", "", "", ""])
    add_table(
        doc,
        ["Client", "Deployment Owner", "Status", "Deployment ID & Time"],
        deploy_rows,
    )

    add_section_title(doc, "Change Details:")
    change_rows = []
    for idx, t in enumerate(tickets, start=1):
        change_rows.append(
            [
                str(idx),
                t.get("key", ""),
                t.get("parent", ""),
                t.get("link", ""),
                t.get("type", ""),
                t.get("client", ""),
                t.get("report", ""),
                t.get("description", ""),
                t.get("test_by", "") if mode == "post" else "",
                t.get("comments", "") if mode == "post" else "",
                t.get("testing_status", "") if mode == "post" else "",
            ]
        )
    change_table = add_table(
        doc,
        [
            "S.No",
            "Ticket#",
            "Parent",
            "Link",
            "Type",
            "Client",
            "Report",
            "Description",
            "Test by",
            "Comments",
            "Testing Status",
        ],
        change_rows or [["", "", "", "", "", "", "", "", "", "", ""]],
    )
    for row_index, ticket in enumerate(tickets, start=1):
        set_cell_hyperlink(change_table.rows[row_index].cells[3], ticket.get("link", ""))

    add_section_title(doc, "Sanity Checklist After Production Deployment:")
    add_bullets(
        doc,
        [
            "Report loads successfully - Yes/No",
            "Paginated report works correctly after applying parameters - Yes/No",
            "Applied changes are reflected as expected - Yes/No",
            "Existing functionalities/changes are not impacted - Yes/No",
        ],
    )

    add_section_title(doc, "Screenshots:")
    # Heading only — team adds screenshots later; do not add placeholder text.

    add_section_title(doc, "Tickets Excluded from Deployment")
    excl_rows = []
    for idx, t in enumerate(excluded, start=1):
        excl_rows.append(
            [
                str(idx),
                t.get("key", ""),
                t.get("parent", ""),
                t.get("link", ""),
                t.get("type", ""),
                t.get("client", ""),
                t.get("report", ""),
                t.get("description", ""),
                t.get("notes", ""),
            ]
        )
    excluded_table = add_table(
        doc,
        [
            "S.No",
            "Ticket#",
            "Parent",
            "Link",
            "Type",
            "Client",
            "Report",
            "Description",
            "Notes",
        ],
        excl_rows or [["", "", "", "", "", "", "", "", ""]],
    )
    for row_index, ticket in enumerate(excluded, start=1):
        set_cell_hyperlink(excluded_table.rows[row_index].cells[3], ticket.get("link", ""))

    return doc


def main():
    parser = argparse.ArgumentParser(description="Generate Fitness BI release notes DOCX")
    parser.add_argument("--input", required=True, help="Path to tickets JSON")
    parser.add_argument("--output", required=True, help="Output DOCX path")
    args = parser.parse_args()

    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document(data)
    doc.save(str(out))
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
