#!/usr/bin/env python3
"""Client-facing Release Notes Word doc, grouped by client.

Shape approved from the 2026.3.09.23 review:
title Release Notes, ticket keys as links, Fix version in References,
no format-reference line and no draft banner.
"""

from __future__ import annotations

import re
from urllib.parse import quote

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

from generate_release_notes import MONTHS

FONT_NAME = "Calibri"
NAVY = RGBColor(0x1F, 0x4E, 0x79)
MENTION = re.compile(r"<custom\b[^>]*>.*?</custom>", re.IGNORECASE | re.DOTALL)
IMAGE = re.compile(r"!\[[^\]]*\]\([^)]*\)")
SKIP_LINE = re.compile(
    r"\b(sow|signing the attached|clear to continue|cc:)\b",
    re.IGNORECASE,
)


def spoken_date(release_date: str, fix_version: str) -> str:
    """23-September-2026 or a fix version -> 23 September 2026."""
    if release_date and "-" in release_date:
        day, month, year = release_date.split("-", 2)
        return f"{int(day)} {month} {year}"
    parts = (fix_version or "").split(".")
    if len(parts) >= 4:
        year, month, day = int(parts[0]), int(parts[2]), int(parts[3])
        return f"{day} {MONTHS[month]} {year}"
    return release_date or fix_version


def issue_lookup(issues: list[dict]) -> dict[str, dict]:
    found = {}
    for issue in issues or []:
        fields = issue.get("fields") or {}
        body = fields.get("description") or ""
        if not isinstance(body, str):
            body = ""
        found[issue.get("key") or ""] = {
            "issue_type": ((fields.get("issuetype") or {}).get("name")) or "",
            "priority": ((fields.get("priority") or {}).get("name")) or "",
            "summary": fields.get("summary") or "",
            "body": body,
        }
    return found


def description_lines(body: str) -> list[str]:
    text = MENTION.sub("", body or "")
    text = IMAGE.sub("", text)
    lines = []
    for raw in text.splitlines():
        line = re.sub(r"\s+", " ", raw).strip(" \t-•*o")
        if not line or SKIP_LINE.search(line):
            continue
        lines.append(line)
    return lines


def implementation_points(ticket: dict, detail: dict) -> list[str]:
    lines = description_lines(detail.get("body") or "")
    if lines:
        return lines
    report = ticket.get("report") or "the report"
    change = ticket.get("description") or ticket.get("highlight") or report
    return [f"{report}: {change}"]


def set_run_font(run, size=12, bold=False, color=None, italic=False):
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_text(doc, text, size=12, bold=False, before=0, after=6, color=None, italic=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.08
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return paragraph


def add_hyperlink(paragraph, text, url, size=12):
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_NAME)
    fonts.set(qn("w:hAnsi"), FONT_NAME)
    props.append(fonts)
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(int(size * 2)))
    props.append(size_el)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    shade.set(qn("w:val"), "clear")
    tc_pr.append(shade)


def write_cell(cell, text, size=9, bold=False, color=None, fill=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text or "")
    set_run_font(run, size=size, bold=bold, color=color)
    if fill:
        shade_cell(cell, fill)


def write_cell_link(cell, text, url, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    add_hyperlink(paragraph, text, url, size=size)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    set_run_font(run, size=12)
    return paragraph


def clients_in_order(tickets: list[dict]) -> list[str]:
    names = []
    for ticket in tickets:
        client = ticket.get("client") or "Unknown"
        if client not in names:
            names.append(client)
    return names


def release_status(payload: dict) -> str:
    return "Deployed" if payload.get("mode") == "post" else "Scheduled"


def fix_version_url(fix_version: str, site: str) -> str:
    jql = quote(f'fixVersion = "{fix_version}" ORDER BY key ASC')
    return f"{site.rstrip('/')}/issues/?jql={jql}"


def opening_sentence(payload: dict) -> str:
    summary = ((payload.get("narratives") or {}).get("deployment_summary")) or ""
    text = summary.replace("This deployment delivers", "This release includes")
    text = text.replace("This deployment contains", "This release contains")
    return text or "This release includes the changes listed below."


def build_client_release_notes(payload: dict, issues: list[dict] | None = None) -> Document:
    tickets = payload.get("tickets") or []
    details = issue_lookup(issues or [])
    fix_version = payload.get("fix_version") or ""
    site = "https://abcfinancial.atlassian.net"
    clients = clients_in_order(tickets)

    doc = Document()
    section = doc.sections[0]
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, margin, Inches(0.8))

    add_text(doc, "Release Notes", size=22, bold=True, after=10, color=NAVY)

    meta = [
        ("Release status", release_status(payload)),
        ("Version", fix_version),
        ("Release date", spoken_date(payload.get("release_date") or "", fix_version)),
        ("Release cycle", "Scheduled production deployment"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.style = "Table Grid"
    for index, (label, value) in enumerate(meta):
        write_cell(table.rows[index].cells[0], label, size=11, bold=True, color=RGBColor(255, 255, 255), fill="1F4E79")
        write_cell(table.rows[index].cells[1], value, size=11)

    add_text(doc, opening_sentence(payload), before=12)

    add_text(doc, "Release summary", size=16, bold=True, before=14, color=NAVY)
    client_phrase = ", ".join(clients) if clients else "none"
    add_bullet(doc, f"Total issues delivered: {len(tickets)}")
    add_bullet(doc, f"Clients: {len(clients)} ({client_phrase})")
    for row in payload.get("category_counts") or []:
        if row.get("count"):
            add_bullet(doc, f"{row.get('label')}: {row['count']}")
    reports = ", ".join(dict.fromkeys(t.get("report") or t.get("description") or t["key"] for t in tickets))
    add_bullet(doc, f"Overall theme: Updates to {reports}" if reports else "Overall theme: Listed client changes")

    add_text(doc, "Included changes", size=16, bold=True, before=14, color=NAVY)
    headers = ["Issue", "Client", "Type", "Summary", "Priority", "Assignee", "Status"]
    changes = doc.add_table(rows=1 + max(len(tickets), 1), cols=7)
    changes.style = "Table Grid"
    for index, header in enumerate(headers):
        write_cell(changes.rows[0].cells[index], header, bold=True, color=RGBColor(255, 255, 255), fill="1F4E79")
    if not tickets:
        write_cell(changes.rows[1].cells[0], "None")
    for index, ticket in enumerate(tickets, start=1):
        detail = details.get(ticket["key"]) or {}
        write_cell_link(changes.rows[index].cells[0], ticket["key"], ticket.get("link") or "", size=9)
        write_cell(changes.rows[index].cells[1], ticket.get("client") or "")
        write_cell(changes.rows[index].cells[2], detail.get("issue_type") or ticket.get("type") or "")
        write_cell(changes.rows[index].cells[3], detail.get("summary") or ticket.get("description") or "")
        write_cell(changes.rows[index].cells[4], detail.get("priority") or "")
        write_cell(changes.rows[index].cells[5], ticket.get("assignee") or "")
        write_cell(changes.rows[index].cells[6], ticket.get("status") or "")

    for client in clients:
        group = [ticket for ticket in tickets if (ticket.get("client") or "Unknown") == client]
        add_text(doc, client, size=18, bold=True, before=16, color=NAVY)
        add_text(doc, "What changed for users", size=14, bold=True, before=8, color=NAVY)
        for ticket in group:
            report = ticket.get("report") or "the listed report"
            change = ticket.get("description") or report
            add_text(
                doc,
                f"The {report} change for {client} is: {change}.",
            )
        impact = "; ".join(
            f"{ticket.get('report') or 'the report'} ({ticket.get('description') or ticket['key']})"
            for ticket in group
        )
        paragraph = add_text(doc, "")
        label = paragraph.add_run("User impact: ")
        set_run_font(label, bold=True)
        detail = paragraph.add_run(
            f"{client} users are affected by {impact}."
        )
        set_run_font(detail)

        add_text(doc, f"Implementation for {client}", size=14, bold=True, before=10, color=NAVY)
        add_text(doc, "Confirmed scope for this client:")
        for ticket in group:
            for point in implementation_points(ticket, details.get(ticket["key"]) or {}):
                add_bullet(doc, point)

        add_text(doc, "Expected impact", size=14, bold=True, before=10, color=NAVY)
        for ticket in group:
            report = ticket.get("report") or "the listed report"
            change = ticket.get("description") or report
            add_bullet(doc, f"{client} can use the updated {report}: {change}.")
        add_bullet(doc, f"Other {client} reports are unchanged unless they are listed in this release.")

        add_text(doc, "Validation", size=14, bold=True, before=10, color=NAVY)
        for ticket in group:
            status = ticket.get("status") or "unknown"
            add_bullet(doc, f"{ticket['key']} Jira status: {status}.")
        if payload.get("mode") == "post":
            add_bullet(doc, "This release is marked deployed.")
        else:
            add_bullet(doc, "Deployment ID and time are recorded after the production deployment.")

    add_text(doc, "References", size=16, bold=True, before=14, color=NAVY)
    for ticket in tickets:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run("Jira: ")
        set_run_font(run)
        add_hyperlink(paragraph, ticket["key"], ticket.get("link") or "", size=12)
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run("Fix version: ")
    set_run_font(run)
    add_hyperlink(paragraph, fix_version, fix_version_url(fix_version, site), size=12)
    return doc
