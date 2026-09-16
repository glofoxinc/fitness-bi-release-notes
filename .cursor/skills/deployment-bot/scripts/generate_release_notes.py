#!/usr/bin/env python3
"""Generate Fitness BI release notes Word doc from tickets JSON."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt


MONTHS = [
    "",
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]


def parse_release_date(fix_version: str) -> str:
    # Version format: YYYY.Quarter.MM.DD  (e.g. 2026.3.08.05 -> 05-August-2026)
    parts = fix_version.strip().split(".")
    if len(parts) >= 4:
        year, month, day = int(parts[0]), int(parts[2]), int(parts[3])
    elif len(parts) == 3:
        # Fallback YYYY.MM.DD
        year, month, day = int(parts[0]), int(parts[1]), int(parts[2])
    else:
        raise ValueError(f"Unexpected fix version format: {fix_version}")
    return f"{day:02d}-{MONTHS[month]}-{year}"


FONT_NAME = "Cambria"
BODY_SIZE = 12
HEADING_SIZE = 14
TITLE_SIZE = 16
TABLE_SIZE = 10


def set_run_font(run, bold=False, size=BODY_SIZE):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FONT_NAME
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), FONT_NAME)
    rFonts.set(qn("w:hAnsi"), FONT_NAME)
    rFonts.set(qn("w:cs"), FONT_NAME)


def add_heading_line(doc: Document, text: str, size=TITLE_SIZE, bold=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, bold=bold, size=size)
    return p


def add_centered_header_block(doc: Document, pairs: list[tuple[str, str]]):
    """Add a centered, borderless header block with aligned colons."""
    table = doc.add_table(rows=len(pairs), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row, (label, value) in zip(table.rows, pairs):
        label_cell, value_cell = row.cells
        label_cell.width = Inches(1.9)
        value_cell.width = Inches(2.7)

        label_paragraph = label_cell.paragraphs[0]
        label_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        label_paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(
            label_paragraph.add_run(f"{label}:"),
            bold=True,
            size=BODY_SIZE,
        )

        value_paragraph = value_cell.paragraphs[0]
        value_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        value_paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(
            value_paragraph.add_run(value or ""),
            bold=False,
            size=BODY_SIZE,
        )
    doc.add_paragraph()
    return table


def add_section_title(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=True, size=HEADING_SIZE)
    return p


def add_body(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=False, size=BODY_SIZE)
    return p


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, bold=False, size=BODY_SIZE)


def set_cell_text(cell, text: str, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text if text is not None else "")
    set_run_font(run, bold=bold, size=TABLE_SIZE)


def set_cell_hyperlink(cell, url: str):
    """Write a real clickable URL instead of plain text."""
    cell.text = ""
    if not url:
        return

    paragraph = cell.paragraphs[0]
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_fonts = OxmlElement("w:rFonts")
    run_fonts.set(qn("w:ascii"), FONT_NAME)
    run_fonts.set(qn("w:hAnsi"), FONT_NAME)
    run_properties.append(run_fonts)

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(TABLE_SIZE * 2))
    run_properties.append(size)

    text = OxmlElement("w:t")
    text.text = url
    run.append(run_properties)
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def shade_header_row(row):
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9E2F3")
        shd.set(qn("w:val"), "clear")
        tc_pr.append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, bold=True)
    shade_header_row(hdr)
    for r_idx, row_vals in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_vals):
            set_cell_text(row.cells[c_idx], val)
    return table


def distinct_clients(tickets: list[dict], typ: str) -> list[str]:
    seen = []
    for t in tickets:
        if t.get("type") != typ:
            continue
        c = t.get("client") or ""
        if c and c not in seen:
            seen.append(c)
    return seen


CATEGORY_LABELS = {
    "new_reports": "New Reports",
    "customizations": "Customizations",
    "looker_exits": "Looker Exits",
    "optimizations": "Optimizations",
}


def category_counts(data: dict) -> list[dict]:
    """Counts come from the builder; recompute only for older tickets JSON."""
    rows = data.get("category_counts")
    if rows:
        return rows
    counts = {key: 0 for key in CATEGORY_LABELS}
    for t in data.get("tickets") or []:
        key = t.get("category")
        if key in counts:
            counts[key] += 1
    return [
        {"key": key, "label": label, "count": counts[key]}
        for key, label in CATEGORY_LABELS.items()
    ]


def change_action(ticket: dict) -> str:
    """Return the change description without repeating the report name."""
    report = (ticket.get("report") or "").strip()
    description = (ticket.get("description") or "").strip()
    if not description or description.lower() == report.lower():
        return ""
    if report and description.lower().startswith(report.lower()):
        return description[len(report) :].strip(" -–—:")
    return description


def build_document(data: dict) -> Document:
    version = data["fix_version"]
    meta = data.get("meta") or {}
    narratives = data.get("narratives") or {}
    tickets = data.get("tickets") or []
    excluded = data.get("excluded") or []
    post = data.get("mode", "pre") == "post"

    release_date = data.get("release_date") or parse_release_date(version)
    analyze_clients = distinct_clients(tickets, "Analyze")
    custom_clients = distinct_clients(tickets, "Custom")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    add_heading_line(doc, "Release Notes", size=18)
    add_centered_header_block(
        doc,
        [
            ("Version", f"v{version}"),
            ("Release Date", release_date),
            ("Development Team", meta.get("development_team", "")),
            ("Release Owner", meta.get("release_owner", "")),
            ("Senior Manager", meta.get("senior_manager", "")),
            ("SVP", meta.get("svp", "")),
            ("Environment", meta.get("environment", "Production")),
        ],
    )

    # One summary section only: brief description followed by the delivery counts.
    add_section_title(doc, "Release Summary:")
    add_body(
        doc,
        narratives.get("release_summary")
        or "This release includes reporting enhancements, custom client updates, and semantic model improvements.",
    )
    size_rows = category_counts(data)
    add_table(
        doc,
        ["Category", "Count"],
        [[row["label"], str(row["count"])] for row in size_rows]
        + [["Total", str(sum(row["count"] for row in size_rows))]],
    )

    add_section_title(doc, "Reports / Dashboards Delivered:")
    delivered = []
    for t in tickets:
        parts = [t.get("client") or "", t.get("report") or ""]
        action = change_action(t)
        if action:
            parts.append(action)
        line = " - ".join(part for part in parts if part)
        if line:
            delivered.append(line)
    add_bullets(doc, delivered or ["(No tickets found)"])

    add_section_title(doc, "Affected Clients:")
    affected = list(analyze_clients) + [c for c in custom_clients if c not in analyze_clients]
    add_bullets(doc, affected or ["None in this release"])

    add_section_title(doc, "Deployment Details:")
    deploy_clients = list(analyze_clients) + [
        client for client in custom_clients if client not in analyze_clients
    ]
    add_table(
        doc,
        ["Client", "Deployment Owner", "Status", "Deployment ID & Time"],
        [[client, "", "", ""] for client in deploy_clients]
        or [["", "", "", ""]],
    )

    add_section_title(doc, "Ticket Details:")
    ticket_rows = [
        [
            ticket.get("key", ""),
            ticket.get("link", ""),
            ticket.get("test_by", "") if post else "",
            ticket.get("deployment_status", "") if post else "",
        ]
        for ticket in tickets
    ]
    ticket_table = add_table(
        doc,
        ["Ticket", "Ticket Link", "Test", "Status"],
        ticket_rows or [["", "", "", ""]],
    )
    for row_index, ticket in enumerate(tickets, start=1):
        set_cell_hyperlink(
            ticket_table.rows[row_index].cells[1],
            ticket.get("link", ""),
        )

    add_section_title(doc, "Testing Evidence:")

    add_section_title(doc, "Tickets Excluded from Deployment:")
    excluded_headers = ["Ticket", "Ticket Link", "Client", "Report", "Notes"]
    excluded_rows = [
        [
            ticket.get("key", ""),
            ticket.get("link", ""),
            ticket.get("client", ""),
            ticket.get("report", ""),
            ticket.get("notes", ""),
        ]
        for ticket in excluded
    ]
    excluded_table = add_table(
        doc,
        excluded_headers,
        excluded_rows or [[""] * len(excluded_headers)],
    )
    for row_index, ticket in enumerate(excluded, start=1):
        set_cell_hyperlink(
            excluded_table.rows[row_index].cells[1],
            ticket.get("link", ""),
        )

    return doc


def main():
    parser = argparse.ArgumentParser(description="Generate Fitness BI release notes DOCX")
    parser.add_argument("--input", required=True, help="Path to tickets JSON")
    parser.add_argument("--output", required=True, help="Output DOCX path")
    args = parser.parse_args()

    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document(data)
    doc.save(str(out))
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
